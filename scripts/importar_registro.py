"""Convierte el «Registro maestro» exportado de Google Sheets (.xlsx) a los CSV de datos/.

Uso (una sola vez, en la migración):
    python3 scripts/importar_registro.py ruta/Registro_maestro.xlsx

- Se saltan Panel, LEEME y Resumen_semanal: ahora los genera scripts/resumen.py.
- Las columnas calculadas con fórmula se eliminan; las calcula scripts/resumen.py.
- Decisiones se escribe como decisiones.md (se lee mejor en el móvil).
"""
import csv
import datetime as dt
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
DATOS = RAIZ / "datos"

HOJAS = {
    "Series_fuerza": "series_fuerza.csv",
    "Sesiones": "sesiones.csv",
    "Carreras": "carreras.csv",
    "Carreras_vueltas": "carreras_vueltas.csv",
    "Otras_actividades": "otras_actividades.csv",
    "Corporal": "corporal.csv",
    "Marcas": "marcas.csv",
    "Historial": "historial.csv",
    "Cambios_perfil": "cambios_perfil.csv",
}


def celda(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def leer_hoja(ws_formulas, ws_valores):
    cab = [c.value for c in ws_formulas[1]]
    ncol = max(i for i, h in enumerate(cab) if h is not None) + 1
    cab = cab[:ncol]
    # Columnas calculadas: alguna celda de datos es una fórmula.
    calculadas = set()
    for fila in ws_formulas.iter_rows(min_row=2, max_col=ncol):
        for i, c in enumerate(fila):
            if isinstance(c.value, str) and c.value.startswith("="):
                calculadas.add(i)
    keep = [i for i in range(ncol) if i not in calculadas]
    filas = []
    for fila in ws_valores.iter_rows(min_row=2, max_col=ncol, values_only=True):
        if all(v in (None, "") for v in fila):
            continue
        filas.append([celda(fila[i]) for i in keep])
    return [cab[i] for i in keep], filas, [cab[i] for i in sorted(calculadas)]


def main(xlsx):
    wf = openpyxl.load_workbook(xlsx)
    wv = openpyxl.load_workbook(xlsx, data_only=True)
    DATOS.mkdir(exist_ok=True)
    for hoja, nombre in HOJAS.items():
        cab, filas, quitadas = leer_hoja(wf[hoja], wv[hoja])
        with open(DATOS / nombre, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(cab)
            w.writerows(filas)
        extra = f" (calculadas, fuera: {', '.join(quitadas)})" if quitadas else ""
        print(f"{nombre}: {len(filas)} filas{extra}")

    cab, filas, _ = leer_hoja(wf["Decisiones"], wv["Decisiones"])
    lineas = ["# Decisiones", "", "Registro fechado de cambios del plan y su motivo. Lo más reciente, arriba.", ""]
    for fecha, decision, motivo in sorted(filas, key=lambda r: r[0], reverse=True):
        lineas.append(f"- **{fecha}** · {decision}" + (f"  \n  *Motivo:* {motivo}" if motivo else ""))
    (RAIZ / "decisiones.md").write_text("\n".join(lineas) + "\n", encoding="utf-8")
    print(f"decisiones.md: {len(filas)} decisiones")


if __name__ == "__main__":
    main(sys.argv[1])
